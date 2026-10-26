import json
import openpyxl
states={'CT','MA','ME','NH','NJ','NY','PA','RI','VT'}
for kind in ('su','mh'):
    wb=openpyxl.load_workbook(f'.local/northeast/samhsa-{kind}.xlsx',read_only=True,data_only=True)
    ws=wb['Facilities List']; it=ws.iter_rows(values_only=True); headers=next(it)
    rows=[dict(zip(headers,r)) for r in it if r[5] in states]
    ws=wb['Service Code Reference']; it=ws.iter_rows(values_only=True); headers=next(it)
    codes=[dict(zip(headers,r)) for r in it if r[2]]
    with open(f'.local/northeast/samhsa-{kind}.json','w',encoding='utf8') as out: json.dump({'rows':rows,'codes':codes},out)
    print(kind,len(rows),{s:sum(r['state']==s for r in rows) for s in sorted(states)})
    print('Payment / facility codes',[(c['service_code'],c['service_name']) for c in codes if c['category_code'] in ('PAY','FT','LANG') or c['service_code'] in ('MC','MD','PI','CASH','SS','PA','SP','ASL','TRICARE')])
